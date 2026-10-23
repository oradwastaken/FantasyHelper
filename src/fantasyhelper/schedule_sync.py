"""Refresh the matchup schedule and remaining-games lookup in Google Sheets."""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from urllib.request import Request, urlopen
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from fantasyhelper.sheets import get_gspread_client


@dataclass(frozen=True)
class ScheduleUpdate:
    matchup_week: int
    start_date: date
    end_date: date
    as_of_date: date
    teams_updated: int


def _excel_schedule(path: str | Path, worksheet_name: str) -> list[list[object]]:
    """Read the fixed populated portion of the local Schedule worksheet."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[worksheet_name]
    except KeyError as error:
        raise ValueError(f"{path} has no {worksheet_name!r} worksheet") from error
    values = list(worksheet.iter_rows(min_row=1, max_row=63, max_col=29, values_only=True))
    if values[0][:3] != ("Week", "Start", "End") or values[30][0] != "Team":
        raise ValueError(f"{path}!{worksheet_name} does not have the expected Schedule layout")
    # gspread sends JSON, which cannot encode Python datetime objects. ISO dates
    # are accepted as dates by Sheets with USER_ENTERED input semantics.
    return [
        [value.isoformat() if isinstance(value, (date, datetime)) else "" if value is None else value for value in row]
        for row in values
    ]


def _matchup_dates(schedule_values: list[list[object]], matchup_week: object) -> tuple[int, date, date]:
    try:
        week = int(matchup_week)
    except (TypeError, ValueError) as error:
        raise ValueError(f"H2H!E37 must contain a numeric matchup week, got {matchup_week!r}") from error
    for row in schedule_values[1:29]:
        if row[0] == week:
            start, end = row[1], row[2]
            if isinstance(start, str): start = date.fromisoformat(start[:10])
            if isinstance(end, str): end = date.fromisoformat(end[:10])
            if isinstance(start, datetime): start = start.date()
            if isinstance(end, datetime): end = end.date()
            if isinstance(start, date) and isinstance(end, date): return week, start, end
    raise ValueError(f"Week {week} is not defined in Schedule!A2:C29")


def _nhl_games_for_week(start_date: date) -> list[dict]:
    """Get the NHL's official game list for the week beginning on ``start_date``."""
    import json
    request = Request(
        f"https://api-web.nhle.com/v1/schedule/{start_date.isoformat()}",
        headers={"User-Agent": "FantasyHelper schedule updater/1.0", "Accept": "application/json"},
    )
    with urlopen(request, timeout=30) as response:  # nosec B310 - fixed NHL API URL
        payload = json.load(response)
    return [game for week in payload.get("gameWeek", []) for game in week.get("games", [])]


def _game_date(game: dict) -> date | None:
    """Return the game's local venue date from either NHL schedule shape."""
    if game.get("gameDate"):
        return date.fromisoformat(game["gameDate"])
    start_time = game.get("startTimeUTC")
    if not start_time:
        return None
    utc_time = datetime.fromisoformat(start_time.replace("Z", "+00:00"))
    return utc_time.astimezone(ZoneInfo(game.get("venueTimezone", "America/Toronto"))).date()


def _nhl_games_for_season(start: date, end: date) -> list[dict]:
    """Fetch each NHL schedule week once, retaining only the sheet's dates."""
    by_id: dict[str, dict] = {}
    current = start
    while current <= end:
        for game in _nhl_games_for_week(current):
            game_date = _game_date(game)
            if game_date is None:
                continue
            if start <= game_date <= end:
                by_id[str(game.get("id", f"{game_date}:{game}"))] = game
        current += timedelta(days=7)
    return list(by_id.values())


def _daily_team_games(games: list[dict], teams: list[str], start: date, end: date) -> tuple[list[date], list[list[int]]]:
    """Build the visible one-column-per-date daily schedule matrix."""
    dates = [start + timedelta(days=offset) for offset in range((end - start).days + 1)]
    date_index = {game_date: index for index, game_date in enumerate(dates)}
    matrix = {team: [0] * len(dates) for team in teams}
    for game in games:
        game_date = _game_date(game)
        if game_date is None:
            continue
        if game_date not in date_index: continue
        for side in ("awayTeam", "homeTeam"):
            abbreviation = str(game.get(side, {}).get("abbrev", "")).strip()
            if abbreviation in matrix: matrix[abbreviation][date_index[game_date]] += 1
    return dates, [matrix[team] for team in teams]


def _column_name(number: int) -> str:
    """Return a one-indexed Google Sheets column name."""
    name = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        name = chr(65 + remainder) + name
    return name


def update_schedule(config: dict, as_of: date | None = None) -> ScheduleUpdate:
    """Copy the local schedule and make H2H W count games remaining including today."""
    settings, sync = config["google_sheets"], config["schedule_sync"]
    spreadsheet_id = settings.get("spreadsheet_id")
    if not spreadsheet_id: raise ValueError("google_sheets.spreadsheet_id is required for update-schedule")
    schedule_values = _excel_schedule(sync["local_workbook"], sync.get("worksheet", "Schedule"))
    client = get_gspread_client(settings["credentials_path"], settings["token_path"])
    spreadsheet = client.open_by_key(spreadsheet_id)
    h2h = spreadsheet.worksheet(sync.get("h2h_worksheet", "H2H"))
    schedule = spreadsheet.worksheet(sync.get("worksheet", "Schedule"))
    player_values = spreadsheet.worksheet(sync.get("player_values_worksheet", "Player Values - Cats"))
    week, start, end = _matchup_dates(schedule_values, h2h.acell(sync.get("matchup_week_cell", "E37")).value)
    as_of = as_of or datetime.now(ZoneInfo(sync.get("timezone", "America/Toronto"))).date()
    if as_of > end: raise ValueError(f"Matchup week {week} ended on {end}; select a current week in H2H!E37")

    schedule.update(schedule_values, "A1:AC63", value_input_option="USER_ENTERED")
    teams = [str(row[0]) for row in schedule_values[31:63] if row[0]]
    season_start = min(_matchup_dates(schedule_values, row[0])[1] for row in schedule_values[1:29])
    season_end = max(_matchup_dates(schedule_values, row[0])[2] for row in schedule_values[1:29])
    daily_dates, daily_games = _daily_team_games(_nhl_games_for_season(season_start, season_end), teams, season_start, season_end)
    daily_start_column = 33  # AG; AD is the live total, AE a spacer, AF a section label.
    daily_end_column = daily_start_column + len(daily_dates) - 1
    daily_start = _column_name(daily_start_column)
    daily_end = _column_name(daily_end_column)
    schedule.update([["Games remaining this week", "", "Daily schedule", *[game_date.isoformat() for game_date in daily_dates]], *[[team, "", "", *games] for team, games in zip(teams, daily_games)]], f"AD31:{daily_end}63", value_input_option="USER_ENTERED")
    remaining_formula = (
        '=SUMPRODUCT(({start}$31:{end}$31>=MAX(TODAY(),INDEX($B$2:$B$29,MATCH(H2H!$E$37,$A$2:$A$29,0))))*'
        '({start}$31:{end}$31<=INDEX($C$2:$C$29,MATCH(H2H!$E$37,$A$2:$A$29,0)))*{start}{row}:{end}{row})'
    )
    schedule.update(
        [[remaining_formula.format(start=daily_start, end=daily_end, row=row)] for row in range(32, 64)],
        "AD32:AD63", value_input_option="USER_ENTERED"
    )
    formula = '=IF($C{row}="","",IFNA(INDEX(Schedule!$AD$32:$AD$63,MATCH($C{row},Schedule!$A$32:$A$63,0)),0))'
    for first, last in ((sync.get("first_roster_row", 3), sync.get("first_roster_last_row", 35)), (sync.get("second_roster_row", 39), sync.get("second_roster_last_row", 71))):
        rows = range(int(first), int(last) + 1)
        h2h.update([[formula.format(row=row)] for row in rows], f"W{rows.start}:W{rows.stop - 1}", value_input_option="USER_ENTERED")
    games_column = sync.get("player_remaining_games_column", "AT")
    player_values.update(
        [["Games remaining this week"], [
            '=ARRAYFORMULA(IF(F3:F="","",IFNA(VLOOKUP(F3:F,Schedule!A32:AD63,30,FALSE),0)))'
        ]],
        f"{games_column}2:{games_column}3", value_input_option="USER_ENTERED"
    )
    return ScheduleUpdate(week, start, end, as_of, len(teams))
